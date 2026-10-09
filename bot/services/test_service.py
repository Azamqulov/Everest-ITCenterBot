import io
import json
import random
import logging
from datetime import datetime
from typing import Optional, Dict, Any, List, Tuple
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from bot.services.db import db

logger = logging.getLogger(__name__)

class TestService:
    @staticmethod
    def generate_test_code() -> str:
        """Generate unique test code, e.g. EV-4821."""
        num = random.randint(1000, 9999)
        return f"EV-{num}"

    @staticmethod
    def parse_answer_key(text: str, question_count: Optional[int] = None) -> Dict[int, str]:
        """
        Parses text like '1a 2c 3b 4d', '1a2b3d', '1A2B3C' or 'abcd' into {1: 'A', 2: 'C', ...}.
        """
        import re
        text = text.strip()
        matches = re.findall(r"(\d+)\s*[:.-]?\s*([a-eA-E])", text)
        result = {}
        if matches:
            for q_str, ans in matches:
                q_num = int(q_str)
                if question_count is None or (1 <= q_num <= question_count):
                    result[q_num] = ans.upper()
        else:
            letters = re.findall(r"[a-eA-E]", text)
            if letters:
                for idx, ans in enumerate(letters, start=1):
                    if question_count is None or idx <= question_count:
                        result[idx] = ans.upper()
        return result

    @staticmethod
    async def create_test(
        teacher_id: int,
        title: str,
        question_count: int,
        answer_key: Dict[int, str],
        subject: Optional[str] = None,
        options_count: int = 4,
        time_limit: int = 30,
        max_attempts: int = 1,
        file_id: Optional[str] = None
    ) -> Dict[str, Any]:
        """Create new test with unique code."""
        code = TestService.generate_test_code()
        while await db.fetch_one("SELECT id FROM tests WHERE code = ?", (code,)):
            code = TestService.generate_test_code()

        await db.execute("""
            INSERT INTO tests (
                teacher_id, title, subject, code, question_count, options_count,
                answer_key, file_id, time_limit, max_attempts, status
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'published')
        """, (
            teacher_id, title, subject, code, question_count, options_count,
            json.dumps(answer_key), file_id, time_limit, max_attempts
        ))

        return await db.fetch_one("SELECT * FROM tests WHERE code = ?", (code,))

    @staticmethod
    async def get_test_by_code(code: str) -> Optional[Dict[str, Any]]:
        return await db.fetch_one("SELECT * FROM tests WHERE code = ? AND status = 'published'", (code.upper().strip(),))

    @staticmethod
    async def get_test_by_id(test_id: int) -> Optional[Dict[str, Any]]:
        return await db.fetch_one("SELECT * FROM tests WHERE id = ?", (test_id,))

    @staticmethod
    async def start_attempt(test_id: int, user_id: int) -> Tuple[Optional[Dict[str, Any]], str]:
        """Start or retrieve an active attempt for user."""
        test = await db.fetch_one("SELECT * FROM tests WHERE id = ?", (test_id,))
        if not test:
            return None, "Test topilmadi."

        # Check existing attempts
        completed_count = await db.fetch_val("""
            SELECT COUNT(*) FROM test_attempts 
            WHERE test_id = ? AND user_id = ? AND status = 'completed'
        """, (test_id, user_id)) or 0

        if completed_count >= test["max_attempts"]:
            return None, f"Siz ushbu testni topshirish imkoniyatidan foydalangansiz (Maksimal: {test['max_attempts']} ta)."

        # Check in_progress attempt
        active_attempt = await db.fetch_one("""
            SELECT * FROM test_attempts 
            WHERE test_id = ? AND user_id = ? AND status = 'in_progress'
        """, (test_id, user_id))

        if active_attempt:
            return active_attempt, "Davom ettirilmoqda"

        # Create new attempt
        await db.execute("""
            INSERT INTO test_attempts (test_id, user_id, total_questions, status)
            VALUES (?, ?, ?, 'in_progress')
        """, (test_id, user_id, test["question_count"]))

        new_att = await db.fetch_one("""
            SELECT * FROM test_attempts 
            WHERE test_id = ? AND user_id = ? AND status = 'in_progress' 
            ORDER BY id DESC LIMIT 1
        """, (test_id, user_id))

        return new_att, "Boshlandi"

    @staticmethod
    async def save_answer(attempt_id: int, question_no: int, chosen_opt: str) -> None:
        """Save student's chosen option in real time."""
        attempt = await db.fetch_one("SELECT * FROM test_attempts WHERE id = ?", (attempt_id,))
        if not attempt:
            return

        test = await db.fetch_one("SELECT answer_key FROM tests WHERE id = ?", (attempt["test_id"],))
        raw_key = test["answer_key"]
        key = json.loads(raw_key) if isinstance(raw_key, str) else raw_key

        correct_ans = str(key.get(str(question_no)) or key.get(question_no) or "").upper()
        is_correct = 1 if chosen_opt.upper() == correct_ans else 0

        # Upsert answer
        existing = await db.fetch_one("SELECT id FROM test_answers WHERE attempt_id = ? AND question_no = ?", (attempt_id, question_no))
        if existing:
            await db.execute("""
                UPDATE test_answers 
                SET chosen = ?, is_correct = ? 
                WHERE attempt_id = ? AND question_no = ?
            """, (chosen_opt.upper(), is_correct, attempt_id, question_no))
        else:
            await db.execute("""
                INSERT INTO test_answers (attempt_id, question_no, chosen, is_correct)
                VALUES (?, ?, ?, ?)
            """, (attempt_id, question_no, chosen_opt.upper(), is_correct))

    @staticmethod
    async def get_user_answers(attempt_id: int) -> Dict[int, str]:
        """Fetch all answered questions for current attempt."""
        rows = await db.fetch_all("SELECT question_no, chosen FROM test_answers WHERE attempt_id = ?", (attempt_id,))
        return {r["question_no"]: r["chosen"] for r in rows}

    @staticmethod
    async def finish_attempt(attempt_id: int) -> Dict[str, Any]:
        """Grade and finish attempt."""
        attempt = await db.fetch_one("SELECT * FROM test_attempts WHERE id = ?", (attempt_id,))
        if not attempt or attempt["status"] == "completed":
            return attempt or {}

        test = await db.fetch_one("SELECT * FROM tests WHERE id = ?", (attempt["test_id"],))
        raw_key = test["answer_key"]
        key = json.loads(raw_key) if isinstance(raw_key, str) else raw_key

        rows = await db.fetch_all("SELECT question_no, chosen FROM test_answers WHERE attempt_id = ?", (attempt_id,))
        answered = {r["question_no"]: r["chosen"] for r in rows}

        correct_count = 0
        total_q = test["question_count"]

        for q in range(1, total_q + 1):
            target = str(key.get(str(q)) or key.get(q) or "").upper()
            if answered.get(q) == target:
                correct_count += 1

        percentage = round((correct_count / total_q) * 100, 1) if total_q > 0 else 0.0
        score = correct_count * 2 # 2 points per correct question

        await db.execute("""
            UPDATE test_attempts 
            SET finished_at = CURRENT_TIMESTAMP,
                score = ?,
                correct_count = ?,
                percentage = ?,
                status = 'completed'
            WHERE id = ?
        """, (score, correct_count, percentage, attempt_id))

        # Award points to user points_ledger
        user_id = attempt["user_id"]
        active_contest = await db.fetch_one("SELECT id FROM contests WHERE status = 'active' ORDER BY id DESC LIMIT 1")
        contest_id = active_contest["id"] if active_contest else None

        await db.execute("""
            INSERT INTO points_ledger (user_id, contest_id, amount, source, reference_id, reason)
            VALUES (?, ?, ?, 'test', ?, ?)
        """, (user_id, contest_id, score, str(test["id"]), f"Test natijasi: {test['title']} ({percentage}%)"))

        updated = await db.fetch_one("SELECT * FROM test_attempts WHERE id = ?", (attempt_id,))
        return {
            **updated,
            "test_title": test["title"],
            "total_questions": total_q,
            "correct_count": correct_count,
            "percentage": percentage,
            "score": score
        }

    @staticmethod
    async def generate_excel_report(test_id: int) -> bytes:
        """Create a beautiful Excel spreadsheet for teacher with test participants and stats."""
        test = await db.fetch_one("SELECT * FROM tests WHERE id = ?", (test_id,))
        attempts = await db.fetch_all("""
            SELECT ta.*, u.full_name, u.phone, u.username
            FROM test_attempts ta
            JOIN users u ON u.telegram_id = ta.user_id
            WHERE ta.test_id = ? AND ta.status = 'completed'
            ORDER BY ta.score DESC, ta.percentage DESC
        """, (test_id,))

        wb = Workbook()
        ws = wb.active
        ws.title = "Natijalar"

        # Styles
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True)
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        # Title
        ws.merge_cells("A1:G1")
        ws["A1"] = f"Everest IT Center — Test Natijalari: {test['title']} ({test['code']})"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(horizontal="center")

        ws["A2"] = f"Savollar: {test['question_count']} ta | Vaqt: {test['time_limit']} daqiqa | Sanasi: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        ws.merge_cells("A2:G2")
        ws["A2"].alignment = Alignment(horizontal="center")

        # Headers
        headers = ["№", "Ism-familiya", "Telefon", "Username", "To'g'ri javoblar", "Foiz", "Ball"]
        ws.append([]) # row 3 blank
        ws.append(headers) # row 4

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=4, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Data rows
        for idx, att in enumerate(attempts, start=1):
            row_data = [
                idx,
                att.get("full_name") or "—",
                att.get("phone") or "—",
                f"@{att.get('username')}" if att.get("username") else "—",
                f"{att['correct_count']} / {att['total_questions']}",
                f"{att['percentage']}%",
                att["score"]
            ]
            ws.append(row_data)
            row_idx = ws.max_row
            for c_idx in range(1, len(row_data) + 1):
                c = ws.cell(row=row_idx, column=c_idx)
                c.border = thin_border
                c.alignment = Alignment(horizontal="center" if c_idx not in [2] else "left")

        # Adjust column widths
        col_widths = [6, 25, 18, 18, 16, 12, 10]
        for col_idx, width in enumerate(col_widths, start=1):
            col_letter = chr(64 + col_idx)
            ws.column_dimensions[col_letter].width = width

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    async def get_test_leaderboard(test_id: int) -> List[Dict[str, Any]]:
        """Fetch all completed test attempts ordered by score, percentage, finished_at."""
        return await db.fetch_all("""
            SELECT ta.*, u.full_name, u.phone, u.username
            FROM test_attempts ta
            JOIN users u ON u.telegram_id = ta.user_id
            WHERE ta.test_id = ? AND ta.status = 'completed'
            ORDER BY ta.score DESC, ta.percentage DESC, ta.finished_at ASC
        """, (test_id,))

    @staticmethod
    async def get_my_created_tests(user_id: int) -> List[Dict[str, Any]]:
        """Fetch all tests created by user with attempt counts."""
        return await db.fetch_all("""
            SELECT t.*, 
                   COUNT(ta.id) as attempts_count,
                   AVG(ta.percentage) as avg_score
            FROM tests t
            LEFT JOIN test_attempts ta ON ta.test_id = t.id AND ta.status = 'completed'
            WHERE t.teacher_id = ?
            GROUP BY t.id
            ORDER BY t.created_at DESC
            LIMIT 20
        """, (user_id,))

    @staticmethod
    async def close_test(test_id: int, user_id: int) -> bool:
        """Close/finish test so no more participants can take it."""
        test = await db.fetch_one("SELECT * FROM tests WHERE id = ?", (test_id,))
        if not test or test["teacher_id"] != user_id:
            return False
        await db.execute("UPDATE tests SET status = 'closed' WHERE id = ?", (test_id,))
        return True

    @staticmethod
    def format_leaderboard(test: Dict[str, Any], attempts: List[Dict[str, Any]]) -> str:
        """Format a human-readable leaderboard ranking text."""
        if not attempts:
            return (
                f"📋 <b>\"{test['title']}\" ({test['code']})</b>\n\n"
                f"Hozircha hech kim ushbu testni topshirmagan.\n"
                f"Ushbu kodni do'stlaringizga ulashing: <code>{test['code']}</code>"
            )

        medals = ["🥇", "🥈", "🥉"]
        lines = [
            f"📊 <b>\"{test['title']}\" ({test['code']}) natijalari:</b>",
            f"👥 Jami qatnashganlar: <b>{len(attempts)} nafar</b>",
            f"❓ Savollar soni: <b>{test['question_count']} ta</b>\n"
        ]

        for idx, att in enumerate(attempts, 1):
            badge = medals[idx - 1] if idx <= 3 else f"{idx}."
            name = att.get("full_name") or "Ishtirokchi"
            correct = att.get("correct_count", 0)
            total = att.get("total_questions", 0)
            pct = att.get("percentage", 0.0)
            pts = att.get("score", 0)
            lines.append(f"{badge} <b>{name}</b> — {correct}/{total} ({pct}%) — <b>+{pts} ball</b>")

        return "\n".join(lines)

test_service = TestService()

